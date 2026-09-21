"""报表服务：把各业务数据组装成 Excel 报表，并支持从导出的报表导回数据。"""

from __future__ import annotations

import datetime as _dt
import os
import re
from dataclasses import dataclass, field

from app.services.app_context import AppContext
from app.utils.excel_export import export_table


class ReportService:
    def __init__(self, ctx: AppContext):
        self.ctx = ctx

    def _dir(self) -> str:
        base = os.path.dirname(self.ctx.db.db_path)
        path = os.path.join(base, "exports")
        os.makedirs(path, exist_ok=True)
        return path

    def _name(self, key: str) -> str:
        return f"{key}_{_dt.date.today().strftime('%Y%m%d')}.xlsx"

    # ---------- 1. 针卡台账 ----------

    def export_cards(self) -> str:
        rows = self.ctx.cards_repo.list_cards()
        data = []
        for r in rows:
            data.append([
                r["name"], r["product_names"], r["status"],
                r["last_touches"], r["last_len"], r["last_update_date"],
                r["in_stock_date"], r["note"],
            ])
        return export_table(
            ["针卡名称", "支持产品", "状态", "当前累计测试量", "当前针长(um)", "最近更新日期", "入库日期", "备注"],
            data, "针卡台账", self._dir(), self._name("针卡台账"),
            col_widths=[14, 26, 8, 16, 13, 13, 12, 24],
        )

    # ---------- 2. 寿命预测 ----------

    def export_predictions(self) -> str:
        data = []
        for p in self.ctx.prediction.predict_all():
            data.append([
                p.card_name, p.product_names, p.status, p.points_count,
                p.last_touches, p.last_len,
                round(p.touches_at_scrap, 0) if p.touches_at_scrap is not None else None,
                round(p.remaining_touches, 0) if p.remaining_touches is not None else None,
                p.rated_touches,
                round(p.achievement_pct, 1) if p.achievement_pct is not None else None,
                p.alert_level,
                round(p.daily_touch_rate, 0) if p.daily_touch_rate else None,
                p.scrap_date,
                {"own": "单卡拟合", "product": "同类型曲线", "none": "数据不足"}.get(p.curve_source, p.curve_source),
            ])
        return export_table(
            ["针卡名称", "支持产品", "状态", "数据点数", "当前累计测试量", "当前针长(um)",
             "预测报废时测试量", "剩余可测次数", "额定寿命", "寿命达成率(%)", "预警级别",
             "消耗速度(次/天)", "预计报废日期", "预测依据"],
            data, "寿命预测", self._dir(), self._name("寿命预测"),
            col_widths=[14, 26, 8, 9, 16, 13, 16, 14, 14, 13, 10, 13, 13, 11],
            notes=[f"报废线 {self.ctx.prediction.scrap_len:g}um；达成率=预测报废时测试量/额定寿命；"
                   f"黄色<{self.ctx.prediction.yellow_pct:g}%，红色<{self.ctx.prediction.red_pct:g}%"],
        )

    # ---------- 3. 需求预估 ----------

    def export_demand(self) -> str:
        data = []
        for d in self.ctx.demand.compute_all():
            arrivals = "；".join(f"第{day}天:{q:.0f}片" for day, q in sorted(d.arrivals_by_day.items()))
            n1_detail = "；".join(f"{k}:{v:.0f}片" for k, v in d.n1_wip_detail.items()) or "无"
            data.append([
                d.product_name, d.touches_per_wafer, d.wph, d.horizon_days,
                round(d.total_wafers), round(d.total_touches),
                d.stock_cards, d.three_card_gap, round(d.capacity_touches),
                d.cards_needed, round(d.n1_wip_total), n1_detail, arrivals,
            ])
        return export_table(
            ["产品", "片耗(次/片)", "WPH", "展望天数", "到达片数(N2)", "测试量需求(次)",
             "在库卡数", "三卡缺口", "现有总余量(次)", "需求针卡张数", "N1在制(片)", "N1分布", "每日到达明细(N2)"],
            data, "需求预估", self._dir(), self._name("需求预估"),
            col_widths=[12, 13, 9, 10, 12, 15, 9, 9, 15, 12, 10, 20, 30],
            notes=["到达量只统计 N2 站点 WIP（按到 CP 时间推算）；N1 站点距 CP 过远只计数量，不参与时间推算；"
                   "三卡缺口按在库卡数判断（在用卡不计入储备）"],
        )

    # ---------- 4. 采购建议与采购单 ----------

    def export_purchase(self) -> str:
        data = []
        for a in self.ctx.purchase.advise_all():
            data.append([
                a.product_name, a.stock_cards, a.three_card_gap, a.in_transit,
                a.horizon_days, round(a.total_touches_needed), round(a.capacity_touches),
                a.shortage_date or "无", a.deadline_date or "—",
                a.urgency_text, a.suggested_qty,
            ])
        headers = ["产品", "在库卡数", "三卡缺口", "在途数量", "展望天数",
                   "测试量需求(次)", "现有总余量(次)", "预计缺口日期", "最晚下单日",
                   "建议", "建议采购数量"]
        pos = self.ctx.purchases_repo.list_all()
        po_rows = [
            [po["id"], po["product_name"], po["qty"], po["order_date"], po["eta_date"],
             po["status"], po["operator"], po["note"]]
            for po in pos
        ]

        from openpyxl import Workbook
        from openpyxl.styles import Alignment
        from openpyxl.utils import get_column_letter
        from app.utils.excel_export import (
            BORDER,
            BODY_FONT,
            HEADER_FILL,
            HEADER_FONT,
            safe_cell,
        )

        wb = Workbook()
        ws = wb.active
        ws.title = "采购建议"
        blocks = [
            ("采购建议", ["产品", "在库卡数", "三卡缺口", "在途数量", "展望天数",
                          "测试量需求(次)", "现有总余量(次)", "预计缺口日期", "最晚下单日",
                          "建议", "建议采购数量"], data),
            ("采购单", ["ID", "产品", "数量", "下单日期", "预计到货", "状态", "操作人", "备注"], po_rows),
        ]
        row = 1
        for title, headers, rows in blocks:
            cell = ws.cell(row=row, column=1, value=title)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            row += 1
            for c, h in enumerate(headers, start=1):
                hc = ws.cell(row=row, column=c, value=h)
                hc.font = HEADER_FONT
                hc.fill = HEADER_FILL
                hc.border = BORDER
                hc.alignment = Alignment(horizontal="center")
            row += 1
            for r in rows:
                for c, v in enumerate(r, start=1):
                    bc = ws.cell(row=row, column=c, value=safe_cell(v) if v is not None else "")
                    bc.font = BODY_FONT
                    bc.border = BORDER
                row += 1
            row += 2
        for c, w in enumerate([12, 10, 10, 10, 10, 16, 16, 14, 14, 24, 14], start=1):
            ws.column_dimensions[get_column_letter(c)].width = w
        path = os.path.join(self._dir(), self._name("采购建议"))
        wb.save(path)
        return path

    # ---------- 报表导入（自动识别导出格式） ----------

    @staticmethod
    def _num(v):
        """宽松数值解析：兼容 1,234,567 / 1.0 / 空白。无法解析返回 None。"""
        if v is None:
            return None
        if isinstance(v, (int, float)):
            return float(v)
        s = str(v).strip().replace(",", "")
        if not s:
            return None
        try:
            return float(s)
        except ValueError:
            return None

    def import_from_excel(self, path: str, operator: str) -> dict:
        """导入之前导出的报表（针卡台账 / 寿命预测），自动识别格式。

        规则：
        - 按表头自动定位数据区并识别报表类型（含"预测报废时测试量"列 → 寿命预测格式，
          含"入库日期"列 → 针卡台账格式）；
        - 支持产品不存在时自动按名称创建（参数留空待补），已存在直接关联；
        - 针卡不存在则创建（名称重复则跳过，只补数据点）；
        - 有累计测试量+针长的行补一条数据点（与最近记录相同时跳过）。
        返回统计 dict：report_type / products_created / cards_created / cards_skipped /
        updates_added / updates_skipped / errors[]。
        """
        from openpyxl import load_workbook

        stats = {
            "report_type": "", "products_created": 0, "cards_created": 0,
            "cards_skipped": 0, "updates_added": 0, "updates_skipped": 0,
            "errors": [],
        }
        wb = load_workbook(path, data_only=True, read_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        wb.close()

        # 定位表头行（前 15 行内包含"针卡名称"）
        header_idx, headers = None, []
        for i, row in enumerate(rows[:15]):
            cells = [str(c).strip() if c is not None else "" for c in row]
            if "针卡名称" in cells:
                header_idx, headers = i, cells
                break
        if header_idx is None:
            raise ValueError("无法识别的报表：表头中未找到【针卡名称】列")

        def col(*names):
            for n in names:
                if n in headers:
                    return headers.index(n)
            return None

        c_name = col("针卡名称")
        c_products = col("支持产品")
        c_status = col("状态")
        c_touches = col("当前累计测试量")
        c_len = col("当前针长(um)")
        c_last_date = col("最近更新日期")
        c_instock = col("入库日期")
        c_note = col("备注")
        c_rated = col("额定寿命")
        if "预测报废时测试量" in headers:
            stats["report_type"] = "寿命预测"
        elif c_instock is not None:
            stats["report_type"] = "针卡台账"
        else:
            stats["report_type"] = "针卡台账（推断）"

        def cell(row, idx):
            if idx is None or idx >= len(row):
                return None
            v = row[idx]
            return str(v).strip() if isinstance(v, str) else v

        today = _dt.date.today().isoformat()
        for rno, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
            name = cell(row, c_name)
            if not name or name in ("—", "-"):
                continue
            name = str(name)
            try:
                pids = []
                raw_products = cell(row, c_products)
                if raw_products:
                    for pname in re.split(r"[、,，/;；]", str(raw_products)):
                        pname = pname.strip()
                        if not pname:
                            continue
                        p = self.ctx.products_repo.get_by_name(pname)
                        if not p:
                            self.ctx.products_repo.create({
                                "name": pname, "touches_per_wafer": 0, "wph": 0,
                                "initial_len_um": 0, "rated_touches": 0,
                                "lead_time_days": 0, "note": "报表导入自动创建，参数待补全",
                            })
                            stats["products_created"] += 1
                            p = self.ctx.products_repo.get_by_name(pname)
                        pids.append(p["id"])
                if not pids:
                    raise ValueError("支持产品为空")

                status = str(cell(row, c_status) or "在库")
                if status not in ("在库", "在用", "报废"):
                    status = "在库"
                rated = self._num(cell(row, c_rated)) or 0
                note = str(cell(row, c_note) or "")
                in_stock = str(cell(row, c_instock) or today)[:10]

                # 先解析数据点再建卡：数值坏的行整行跳过并报告，不产生半截数据
                touches = self._num(cell(row, c_touches))
                needle_len = self._num(cell(row, c_len))
                if (touches is None) != (needle_len is None) or (
                    touches is not None and (touches <= 0 or needle_len <= 0)
                ):
                    raise ValueError("累计测试量/针长不是有效数值（或为非正数）")

                card = self.ctx.cards_repo.get_by_name(name)
                if card is None:
                    card_id = self.ctx.cards.create_card(
                        operator=operator, name=name, product_ids=pids, status=status,
                        rated_touches=rated, in_stock_date=in_stock, note=note,
                    )
                    stats["cards_created"] += 1
                else:
                    card_id = card["id"]
                    stats["cards_skipped"] += 1

                if touches is None or needle_len is None:
                    stats["updates_skipped"] += 1
                    continue
                last = self.ctx.cards_repo.latest_update(card_id)
                if last and abs(float(last["cum_touches"]) - touches) < 0.5 \
                        and abs(float(last["needle_len"]) - needle_len) < 0.05:
                    stats["updates_skipped"] += 1
                    continue
                udate = str(cell(row, c_last_date) or today)[:10]
                try:
                    _dt.datetime.strptime(udate, "%Y-%m-%d")
                except ValueError:
                    udate = today
                self.ctx.cards.add_update(
                    card_id=card_id, update_date=udate, cum_touches=touches,
                    needle_len=needle_len, operator=operator, note="报表导入",
                )
                stats["updates_added"] += 1
            except Exception as e:  # noqa: BLE001
                stats["errors"].append(f"第 {rno} 行【{name}】：{e}")
        self.ctx.audit_repo.log(
            operator, "报表导入", "导入完成",
            f"类型 {stats['report_type']}：新卡 {stats['cards_created']}，"
            f"跳过 {stats['cards_skipped']}，数据点 {stats['updates_added']}，"
            f"失败 {len(stats['errors'])}",
        )
        return stats
