"""Excel 导出工具（openpyxl），统一表头样式。

安全：用户可自由输入的字段（针卡名、备注等）若以 = + - @ 开头，写入 xlsx 会被
Excel 当作公式执行（公式注入）。导出前统一在前面补一个空格使其变成纯文本。
"""

from __future__ import annotations

import os
from typing import Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

_FORMULA_PREFIXES = ("=", "+", "-", "@")
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name="微软雅黑", size=10)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def safe_cell(value):
    """阻断 Excel 公式注入：以 = + - @ 开头的字符串前置空格转为纯文本。"""
    if isinstance(value, str) and value.startswith(_FORMULA_PREFIXES):
        return " " + value
    return value


def export_table(
    headers: Sequence[str],
    rows: Sequence[Sequence[object]],
    title: str,
    default_dir: str,
    suggested_name: str,
    col_widths: Sequence[int] | None = None,
    notes: Sequence[str] = (),
) -> str:
    """导出单表 Excel，返回保存路径。rows 中的 None 输出为空。"""
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31] or "Sheet1"

    r0 = 1
    if notes:
        for i, note in enumerate(notes):
            cell = ws.cell(row=1 + i, column=1, value=note)
            cell.font = Font(name="微软雅黑", size=9, color="808080")
        r0 = 1 + len(notes) + 1

    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=r0, column=c, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r, row in enumerate(rows, start=r0 + 1):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=safe_cell(v) if v is not None else "")
            cell.font = BODY_FONT
            cell.border = BORDER

    for c in range(1, len(headers) + 1):
        width = col_widths[c - 1] if col_widths and c <= len(col_widths) else 14
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.freeze_panes = ws.cell(row=r0 + 1, column=1)

    os.makedirs(default_dir, exist_ok=True)
    path = os.path.join(default_dir, suggested_name)
    wb.save(path)
    return path
