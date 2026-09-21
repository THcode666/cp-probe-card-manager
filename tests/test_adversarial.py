"""上线前对抗性测试：恶意用户 / 极端数据 / 无权限用户 三类攻击面的回归。"""

from __future__ import annotations

import math
import os
import sys

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app.services.app_context import AppContext  # noqa: E402


@pytest.fixture()
def ctx(tmp_path):
    c = AppContext(str(tmp_path / "adv.db"))
    c.current_user = {"id": 1, "username": "admin", "role": "管理员"}
    return c


def _mk_product(ctx, name="0787"):
    return ctx.products_repo.create({"name": name, "touches_per_wafer": 17146, "wph": 60,
                                     "initial_len_um": 30.0, "rated_touches": 4800000,
                                     "lead_time_days": 30, "note": ""})


# ============ 一、恶意用户 ============

def test_excel_formula_injection_blocked(ctx, tmp_path):
    """恶意用户：针卡名塞 Excel 公式 → 导出必须被消毒为纯文本。"""
    pid = _mk_product(ctx)
    evil = '=HYPERLINK("http://evil.example","点我领奖")'
    cid = ctx.cards.create_card("admin", evil, [pid])
    ctx.cards.add_update(card_id=cid, update_date="2026-09-12", cum_touches=1000,
                         needle_len=29.5, operator="admin")
    from app.services.report_service import ReportService
    path = ReportService(ctx).export_cards()
    from openpyxl import load_workbook
    ws = load_workbook(path).active
    cell = ws.cell(row=2, column=1).value
    assert not str(cell).startswith("="), "导出单元格以 = 开头会被 Excel 当公式执行"
    assert "HYPERLINK" in str(cell), "内容应保留（仅消毒，不丢数据）"


def test_login_lockout_after_5_failures(ctx):
    """恶意用户：连续猜错 5 次密码 → 锁定 5 分钟，正确密码也进不来。"""
    for i in range(5):
        with pytest.raises(ValueError):
            ctx.auth.login("admin", f"wrong-{i}")
    with pytest.raises(ValueError, match="锁定"):
        ctx.auth.login("admin", "Admin@12345")  # 即使密码正确也被拒
    # 解锁（模拟 5 分钟后）→ 正常登录
    ctx.db.execute("UPDATE users SET locked_until = '2000-01-01 00:00:00' WHERE username = 'admin'")
    user = ctx.auth.login("admin", "Admin@12345")
    assert user["username"] == "admin"


def test_login_success_resets_failure_counter(ctx):
    with pytest.raises(ValueError):
        ctx.auth.login("admin", "wrong")
    with pytest.raises(ValueError):
        ctx.auth.login("admin", "wrong")
    ctx.auth.login("admin", "Admin@12345")  # 成功清零
    u = ctx.users_repo.get_by_username("admin")
    assert u["failed_attempts"] == 0 and not u["locked_until"]


def test_rich_text_in_card_name_does_not_break_ui(ctx):
    """恶意用户：卡名带 HTML → 详情页只做转义展示，不影响结构（预测对象字段可取）。"""
    pid = _mk_product(ctx)
    evil = "RR<b>坏</b>"
    cid = ctx.cards.create_card("admin", evil, [pid])
    pred = ctx.prediction.predict_card(cid)
    assert pred.card_name == evil  # 数据原样存储；界面层负责转义
    import html
    assert html.escape(evil) in html.escape(pred.card_name)


# ============ 二、无权限用户 ============

def test_engineer_cannot_call_admin_apis(ctx):
    """无权限用户：绕过界面直接调管理员接口 → 服务层拦截。"""
    pid = _mk_product(ctx)
    cid = ctx.cards.create_card("admin", "RR-01", [pid])
    ctx.auth.register("engineer1", "pass123", "change-me-activation")
    eng_id = ctx.users_repo.get_by_username("engineer1")["id"]
    ne_id = ctx.users_repo.get_by_username("admin")["id"]

    with pytest.raises(ValueError, match="管理员权限"):
        ctx.auth.reveal_password("engineer1", ne_id)
    with pytest.raises(ValueError, match="管理员权限"):
        ctx.auth.delete_user("engineer1", ne_id)
    with pytest.raises(ValueError, match="管理员权限"):
        ctx.auth.create_user("engineer1", "x", "pass123", "工程师")
    with pytest.raises(ValueError, match="管理员权限"):
        ctx.auth.reset_password("engineer1", eng_id, "pass123")
    with pytest.raises(ValueError, match="管理员权限"):
        ctx.auth.set_role("engineer1", eng_id, "管理员")
    with pytest.raises(ValueError, match="管理员权限"):
        ctx.cards.delete_card("engineer1", cid)
    # 被拒操作都有审计记录
    logs = ctx.audit_repo.list_recent()
    assert any("越权尝试" in (l["action"] + l["detail"]) for l in logs)


def test_inactive_admin_cannot_use_admin_apis(ctx):
    """无权限用户：被停用的管理员接口调用同样被拒。"""
    admin = ctx.users_repo.get_by_username("admin")
    ctx.users_repo.set_active(admin["id"], False)
    with pytest.raises(ValueError, match="管理员权限"):
        ctx.auth.delete_user("admin", 1)


def test_self_role_downgrade_blocked(ctx):
    with pytest.raises(ValueError, match="降级"):
        ctx.auth.set_role("admin", ctx.users_repo.get_by_username("admin")["id"], "工程师")


# ============ 三、极端数据制造者 ============

def test_extreme_update_values_rejected(ctx):
    pid = _mk_product(ctx)
    cid = ctx.cards.create_card("admin", "RR-01", [pid])
    _add = lambda t, l, d="2026-09-12": ctx.cards.add_update(
        card_id=cid, update_date=d, cum_touches=t, needle_len=l, operator="admin")
    with pytest.raises(ValueError, match="范围"):
        _add(1e13, 29.0)                       # 天文数字
    with pytest.raises(ValueError, match="范围"):
        _add(-5, 29.0)                         # 负数
    with pytest.raises(ValueError, match="范围"):
        _add(1000, 5000)                       # 针长超物理极限
    with pytest.raises(ValueError, match="有限数值"):
        _add(float("nan"), 29.0)               # NaN
    with pytest.raises(ValueError, match="有限数值"):
        _add(float("inf"), 29.0)               # Inf
    with pytest.raises(ValueError, match="日期"):
        _add(100, 29.0, "2026-13-40")          # 非法日期
    with pytest.raises(ValueError, match="备注"):
        ctx.cards.add_update(card_id=cid, update_date="2026-09-12", cum_touches=100,
                             needle_len=29.0, operator="admin", note="长" * 501)


def test_card_and_product_name_length_capped(ctx):
    pid = _mk_product(ctx)
    with pytest.raises(ValueError, match="过长"):
        ctx.cards.create_card("admin", "卡" * 100, [pid])
    with pytest.raises(ValueError, match="过长"):
        ctx.products_repo.create({"name": "品" * 100, "touches_per_wafer": 1, "wph": 1,
                                  "initial_len_um": 30, "rated_touches": 0, "lead_time_days": 1,
                                  "note": ""})


def test_duplicate_station_friendly_error(ctx):
    ctx.planning_repo.create_station("LOT", 3)
    with pytest.raises(ValueError, match="已存在"):
        ctx.planning_repo.create_station("LOT", 5)
    with pytest.raises(ValueError, match="过长"):
        ctx.planning_repo.create_station("站" * 100, 1)


def test_wip_qty_bounds(ctx):
    pid = _mk_product(ctx)
    ctx.planning_repo.create_station("LOT", 3)
    sid = ctx.planning_repo.list_stations()[0]["id"]
    with pytest.raises(ValueError, match="WIP"):
        ctx.planning_repo.set_wip(pid, sid, -10, "tester")
    with pytest.raises(ValueError, match="WIP"):
        ctx.planning_repo.set_wip(pid, sid, 2_000_000_001, "tester")
    ctx.planning_repo.set_wip(pid, sid, 100, "tester")  # 正常值可存


def test_settings_bounds_reject_extremes(ctx):
    """极端数据：把报废线/阈值改成荒谬值会被拒绝，预测体系不被污染。"""
    with pytest.raises(ValueError, match="取值须在"):
        ctx.settings_repo.set("scrap_needle_len_um", "-5", "admin")
    with pytest.raises(ValueError, match="取值须在"):
        ctx.settings_repo.set("warn_red_pct", "1e9", "admin")
    with pytest.raises(ValueError, match="取值须在"):
        ctx.settings_repo.set("demand_horizon_days", "0", "admin")
    with pytest.raises(ValueError, match="数字"):
        ctx.settings_repo.set("min_cards_per_product", "三张", "admin")
    # 合法值仍可写入
    ctx.settings_repo.set("scrap_needle_len_um", "20", "admin")
    assert ctx.settings_repo.get_float("scrap_needle_len_um") == 20


def test_prediction_survives_hostile_curve(ctx):
    """极端数据：数据点恶意乱填（针长不降反升/恒定）→ 预测返回"无法预测"而不是崩溃。"""
    pid = _mk_product(ctx)
    cid = ctx.cards.create_card("admin", "RR-01", [pid])
    for d, t, l in [("2026-08-01", 0, 30.0), ("2026-08-11", 100, 31.0), ("2026-08-21", 200, 32.0)]:
        ctx.cards.add_update(card_id=cid, update_date=d, cum_touches=t, needle_len=l, operator="admin")
    pred = ctx.prediction.predict_card(cid)
    assert pred.touches_at_scrap is None  # 针长上升 → 不外推，不崩溃
